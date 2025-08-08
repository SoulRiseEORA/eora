import os
import re
import asyncio
from datetime import datetime
from pathlib import Path
from typing import List, Optional

# 파일별 텍스트 추출용 패키지
try:
    import docx
except ImportError:
    docx = None
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None
try:
    import pandas as pd
except ImportError:
    pd = None
try:
    import openpyxl
except ImportError:
    openpyxl = None

# 토큰화
try:
    import tiktoken
    enc = tiktoken.get_encoding("cl100k_base")
    def count_tokens(text):
        return len(enc.encode(text))
    def split_by_tokens(text, max_tokens=1500):
        tokens = enc.encode(text)
        chunks = []
        for i in range(0, len(tokens), max_tokens):
            chunk = enc.decode(tokens[i:i+max_tokens])
            chunks.append(chunk)
        return chunks
except ImportError:
    def count_tokens(text):
        return len(text.split())
    def split_by_tokens(text, max_tokens=1500):
        words = text.split()
        chunks = []
        for i in range(0, len(words), max_tokens):
            chunk = " ".join(words[i:i+max_tokens])
            chunks.append(chunk)
        return chunks

# 텍스트 추출 함수

def extract_text_from_file(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()
    if ext == '.txt' or ext == '.md' or ext == '.py':
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    elif ext == '.docx':
        if docx is None:
            raise ImportError('python-docx 패키지가 필요합니다.')
        doc = docx.Document(file_path)
        return '\n'.join([p.text for p in doc.paragraphs])
    elif ext == '.pdf':
        if PyPDF2 is None:
            raise ImportError('PyPDF2 패키지가 필요합니다.')
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            text = ''
            for page in reader.pages:
                text += page.extract_text() + '\n'
            return text
    elif ext in ['.xlsx', '.xls']:
        if pd is not None:
            df = pd.read_excel(file_path, dtype=str)
            return '\n'.join(df.astype(str).apply(lambda row: ' | '.join(row), axis=1))
        elif openpyxl is not None and ext == '.xlsx':
            wb = openpyxl.load_workbook(file_path)
            text = ''
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    text += ' | '.join([str(cell) if cell is not None else '' for cell in row]) + '\n'
            return text
        else:
            raise ImportError('엑셀 파일 처리를 위해 pandas 또는 openpyxl 패키지가 필요합니다.')
    else:
        raise ValueError('지원하지 않는 파일 형식: ' + ext)

# Aura Memory 저장 함수
async def save_chunks_to_aura_memory(chunks: List[str], file_name: str, user: str = "system"):
    """
    각 청크를 아우라 메모리 시스템에 저장
    """
    from aura_system.memory_manager import get_memory_manager
    memory_manager = await get_memory_manager()
    for idx, chunk in enumerate(chunks):
        metadata = {
            "type": "file_chunk",
            "file_name": file_name,
            "chunk_index": idx,
            "user": user,
            "timestamp": datetime.utcnow().isoformat(),
        }
        await memory_manager.store_memory(content=chunk, metadata=metadata)

# 전체 파이프라인: 파일 → 텍스트 추출 → 분할 → 저장
async def learn_file_to_aura_memory(file_path: str, user: str = "system", max_tokens: int = 1500):
    file_name = os.path.basename(file_path)
    text = extract_text_from_file(file_path)
    chunks = split_by_tokens(text, max_tokens=max_tokens)
    await save_chunks_to_aura_memory(chunks, file_name, user)
    return len(chunks)

# 대화파일(턴별) 분할 및 저장
async def learn_dialog_file_to_aura_memory(file_path: str, user: str = "system"):
    file_name = os.path.basename(file_path)
    text = extract_text_from_file(file_path)
    # turn: 사용자: ...\nGPT: ...\n 또는 Q: ...\nA: ...\n 등 패턴 분할
    turns = re.split(r'(?:^|\n)(?:Q:|사용자:|User:|\[USER\])', text)
    turn_count = 0
    for idx, turn in enumerate(turns):
        if not turn.strip():
            continue
        # 답변 추출
        m = re.search(r'(?:A:|GPT:|Assistant:|\[GPT\])(.+)', turn, re.DOTALL)
        if m:
            user_input = turn.split(m.group(0))[0].strip()
            gpt_response = m.group(1).strip()
        else:
            user_input = turn.strip()
            gpt_response = ""
        from aura_system.memory_structurer import MemoryAtom
        from aura_system.memory_manager import get_memory_manager
        memory_manager = await get_memory_manager()
        metadata = {
            "type": "conversation",
            "file_name": file_name,
            "turn_index": idx,
            "user": user,
            "timestamp": datetime.utcnow().isoformat(),
        }
        atom = MemoryAtom(content=user_input + "\n" + gpt_response, metadata=metadata)
        await memory_manager.store_memory(content=atom.content, metadata=atom.metadata)
        turn_count += 1
    return turn_count 