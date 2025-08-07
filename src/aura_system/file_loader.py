import sys, os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "aura_system")))
import os
import json
from typing import List, Dict
from datetime import datetime
from aura_system.embedding_engine import embed_text
import asyncio
from aura_system.memory_manager import get_memory_manager
from aura_system.meta_store import get_meta_store
from pathlib import Path
try:
    import docx
except ImportError:
    docx = None
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None
try:
    import openpyxl
except ImportError:
    openpyxl = None
try:
    import pandas as pd
except ImportError:
    pd = None


def split_text_into_chunks(text: str, max_length: int = 1000) -> List[str]:
    lines = text.split('\n')
    chunks = []
    chunk = ""
    for line in lines:
        if len(chunk) + len(line) < max_length:
            chunk += line + "\n"
        else:
            chunks.append(chunk.strip())
            chunk = line + "\n"
    if chunk:
        chunks.append(chunk.strip())
    return chunks


async def load_file_and_store_memory(file_path: str, file_name: str = None):
    if not file_name:
        file_name = os.path.basename(file_path)

    ext = Path(file_path).suffix.lower()
    text = None
    if ext == '.txt':
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
    elif ext == '.docx':
        if docx is None:
            raise ImportError('python-docx 패키지가 설치되어 있지 않습니다.')
        doc = docx.Document(file_path)
        text = '\n'.join([p.text for p in doc.paragraphs])
    elif ext == '.pdf':
        if PyPDF2 is None:
            raise ImportError('PyPDF2 패키지가 설치되어 있지 않습니다.')
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            text = ''
            for page in reader.pages:
                text += page.extract_text() + '\n'
    elif ext in ['.xlsx', '.xls']:
        if pd is not None:
            df = pd.read_excel(file_path, dtype=str)
            text = '\n'.join(df.astype(str).apply(lambda row: ' | '.join(row), axis=1))
        elif openpyxl is not None and ext == '.xlsx':
            wb = openpyxl.load_workbook(file_path)
            text = ''
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    text += ' | '.join([str(cell) if cell is not None else '' for cell in row]) + '\n'
        else:
            raise ImportError('엑셀 파일 처리를 위해 pandas 또는 openpyxl 패키지가 필요합니다.')
    else:
        raise ValueError('지원하지 않는 파일 형식입니다: ' + ext)

    print(f"파일 '{file_path}'에서 추출한 전체 텍스트 길이: {len(text)}자")
    chunks = split_text_into_chunks(text)
    print(f"분할된 청크 개수: {len(chunks)}")
    memory_manager = await get_memory_manager()
    meta_store = await get_meta_store()
    for idx, chunk in enumerate(chunks):
        print(f"청크 {idx} 타입: {type(chunk)}")
        if not isinstance(chunk, str):
            print(f"청크 {idx} 타입이 {type(chunk)}이므로 문자열로 변환합니다.")
            chunk = " ".join([str(c) for c in chunk])
        token_count = len(chunk.split())
        print(f"청크 {idx} 토큰 수: {token_count}")
        emb = await embed_text(chunk)
        memory_metadata = {
            "type": "file_chunk",
            "file_name": file_name,
            "chunk_index": idx,
            "tags": extract_tags(chunk),
            "summary_prompt": summarize_text(chunk),
            "resonance_score": estimate_resonance(chunk),
            "timestamp": datetime.now().isoformat()
        }
        # 메모리 저장
        ok = await memory_manager.store_memory(content=chunk, metadata=memory_metadata)
        if not ok:
            print(f"청크 {idx} 저장 실패: 메타데이터={memory_metadata} (원인: content가 비었거나, 중복, DB 연결, 임베딩 등)")
        else:
            print(f"청크 {idx} 저장 성공: 메타데이터={memory_metadata}")
        # 메타데이터 저장 (memory_id는 store_memory에서 반환받아야 정확, 여기선 생략 또는 임시)
        await meta_store.store_metadata(
            memory_id=f"{file_name}_chunk_{idx}_{int(datetime.now().timestamp())}",
            metadata=memory_metadata
        )


def extract_tags(text: str) -> List[str]:
    words = text.lower().split()
    return list(set(words))


def summarize_text(text: str) -> str:
    return text[:80].replace('\n', ' ') + "..."


def estimate_resonance(text: str) -> int:
    return min(100, 60 + len(text) % 40)


if __name__ == "__main__":
    test_path = "./docs/test_article.txt"
    asyncio.run(load_file_and_store_memory(test_path))
    print("✅ 파일 학습 및 기억 저장 완료")